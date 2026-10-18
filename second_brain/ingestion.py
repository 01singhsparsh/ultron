import os, sys, json, datetime, hashlib

BASE_DIR      = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KNOWLEDGE_DIR = os.path.join(BASE_DIR, "data", "knowledge")
INGESTED_FILE = os.path.join(BASE_DIR, "data", "ingested_files.json")
os.makedirs(KNOWLEDGE_DIR, exist_ok=True)

SUPPORTED = ['.pdf','.docx','.pptx','.txt','.md','.py','.csv']

def _file_hash(path):
    with open(path,'rb') as f: return hashlib.md5(f.read()).hexdigest()

def _load_ingested():
    try:
        if os.path.exists(INGESTED_FILE):
            with open(INGESTED_FILE) as f: return json.load(f)
    except: pass
    return {}

def _save_ingested(data):
    with open(INGESTED_FILE,'w') as f: json.dump(data, f, indent=2)

def extract_text_from_pdf(path):
    try:
        import pypdf
        reader = pypdf.PdfReader(path)
        return "".join([page.extract_text() or "" for page in reader.pages]).strip()
    except Exception as e: print(f"[INGEST] PDF error: {e}"); return ""

def extract_text_from_docx(path):
    try:
        import docx
        doc = docx.Document(path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()]).strip()
    except Exception as e: print(f"[INGEST] DOCX error: {e}"); return ""

def extract_text_from_pptx(path):
    try:
        from pptx import Presentation
        prs = Presentation(path); text = ""
        for i, slide in enumerate(prs.slides):
            text += f"\n--- Slide {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape,"text") and shape.text.strip():
                    text += shape.text.strip() + "\n"
        return text.strip()
    except Exception as e: print(f"[INGEST] PPTX error: {e}"); return ""

def extract_text_from_txt(path):
    try:
        with open(path,'r',encoding='utf-8',errors='ignore') as f: return f.read().strip()
    except Exception as e: print(f"[INGEST] TXT error: {e}"); return ""

def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pdf':                          return extract_text_from_pdf(path)
    elif ext == '.docx':                       return extract_text_from_docx(path)
    elif ext == '.pptx':                       return extract_text_from_pptx(path)
    elif ext in ['.txt','.md','.py','.csv']:  return extract_text_from_txt(path)
    return ""

def ingest_file(path, category="general"):
    ingested  = _load_ingested()
    file_hash = _file_hash(path)
    if path in ingested and ingested[path].get("hash") == file_hash:
        print(f"[INGEST] Already ingested: {os.path.basename(path)}")
        return ingested[path]
    print(f"[INGEST] Processing: {os.path.basename(path)}")
    text = extract_text(path)
    if not text or len(text) < 50:
        print(f"[INGEST] No text extracted from: {path}"); return {}
    try:
        from second_brain.vectorstore import add_document
        chunks_added = add_document(text=text, metadata={
            "source":   os.path.basename(path), "path": path,
            "category": category, "type": os.path.splitext(path)[1].lower(),
            "ingested": str(datetime.datetime.now()),
        })
        record = {"hash": file_hash, "source": os.path.basename(path),
                  "category": category, "chunks": chunks_added,
                  "size": len(text), "ingested_at": str(datetime.datetime.now())}
        ingested[path] = record; _save_ingested(ingested)
        print(f"[INGEST] ✓ {os.path.basename(path)} — {chunks_added} chunks stored")
        return record
    except Exception as e:
        print(f"[INGEST] Vector store error: {e}"); return {}

def ingest_folder(folder=None, category="general"):
    folder = folder or KNOWLEDGE_DIR; count = 0
    for fname in os.listdir(folder):
        if os.path.splitext(fname)[1].lower() in SUPPORTED:
            result = ingest_file(os.path.join(folder, fname), category)
            if result: count += 1
    return count

def get_ingested_files():
    ingested = _load_ingested()
    return [{"name": v.get("source",""), "category": v.get("category",""),
             "chunks": v.get("chunks",0), "size": v.get("size",0),
             "date": v.get("ingested_at","")[:10]} for v in ingested.values()]

def get_knowledge_dir(): return KNOWLEDGE_DIR